"""本地用例表读取：xlsx / csv。"""

from __future__ import annotations

import csv
from pathlib import Path

from work_agent.tools.models import SheetTable


class LocalCaseSheetTool:
    """确定性读表：不猜测列语义，全部当字符串。"""

    def read(self, path: str) -> SheetTable:
        """
        读取本地 xlsx/csv 用例表。

        参数:
            path: 文件路径。

        返回:
            SheetTable；文件不存在或不支持格式时抛错。
        """
        p = Path(path).expanduser()
        if not p.is_file():
            raise FileNotFoundError(f"用例表不存在: {p}")

        suffix = p.suffix.lower()
        if suffix == ".csv":
            headers, rows = self._read_csv(p)
        elif suffix in {".xlsx", ".xlsm"}:
            headers, rows = self._read_xlsx(p)
        else:
            raise ValueError(f"不支持的表格格式: {suffix}（请用 .xlsx 或 .csv）")

        return SheetTable(headers=headers, rows=rows, path=str(p.resolve()))

    @staticmethod
    def _read_csv(path: Path) -> tuple[list[str], list[list[str]]]:
        """读 CSV 为表头 + 行。"""
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            reader = csv.reader(f)
            all_rows = [[(c or "").strip() for c in row] for row in reader]
        if not all_rows:
            return [], []
        headers = all_rows[0]
        rows = all_rows[1:]
        width = len(headers)
        norm_rows = [r + [""] * (width - len(r)) if len(r) < width else r[:width] for r in rows]
        return headers, norm_rows

    @staticmethod
    def _read_xlsx(path: Path) -> tuple[list[str], list[list[str]]]:
        """读 xlsx 活动表为表头 + 行（需 openpyxl）。"""
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            raise ImportError("读取 xlsx 需要安装 openpyxl") from exc

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        raw: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            raw.append(["" if c is None else str(c).strip() for c in row])
        wb.close()
        if not raw:
            return [], []
        headers = raw[0]
        # 去掉表头右侧全空列
        while headers and headers[-1] == "":
            headers = headers[:-1]
        width = len(headers)
        rows = []
        for r in raw[1:]:
            cells = (r + [""] * width)[:width]
            if any(cells):
                rows.append(cells)
        return headers, rows
